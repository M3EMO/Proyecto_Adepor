from openpyxl import Workbook

from src.persistencia.excel_hoja_backtest import poblar_backtest


def _fila(fecha):
    # 26 columnas base que espera poblar_backtest (ver motor_sincronizador._cargar_partidos)
    return ['id', fecha, 'A', 'B', 'Argentina', 0.5, 0.25, 0.25, 0.5, 0.5,
            '[GANADA] LOCAL', '', 10.0, 0.0, 2.0, 3.0, 4.0, 1.9, 1.9,
            'Liquidado', 1, 0, 0.1, 'NO', '', 0.0]


def test_aporte_posterior_al_ultimo_partido_entra_en_equity():
    wb = Workbook()
    poblar_backtest(wb, [_fila('2026-04-01 20:00')], 1000.0,
                    aportes=[('2026-05-01', 500.0)])
    ws = wb['Backtest'] if 'Backtest' in wb.sheetnames else wb.worksheets[-1]
    formulas = [str(c.value) for row in ws.iter_rows(min_row=2, max_row=2) for c in row]
    assert any('500' in f for f in formulas), formulas
